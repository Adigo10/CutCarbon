"""Primary-data coverage % and user-declared boundary exclusions.

Two pieces of "how this number was built" provenance:

* `coverage_pct` — the share of the footprint that comes from categories flagged
  "actual" in `assumptions.category_data_quality`. Derived at read time so stored
  scenarios pick it up without a recalculation.
* `exclusions` — free text the organizer writes to declare what is outside the
  reporting boundary. It rides in `input_payload` and must reach every export;
  when it is empty the reports say "None declared" rather than going silent.
"""

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError

from app.models.database import ScenarioDB
from app.models.schemas import (
    ComplianceReport,
    EventScenarioInput,
    OffsetPortfolioSummary,
    ScenarioReportPayload,
    ScopeBreakdown,
)
from app.routers.exports import methodology_block
from app.services.scenario_serializer import primary_data_coverage_pct, serialize_scenario

from helpers import create_scenario, register_user

EXCLUSIONS_LABEL = "Boundary exclusions (user-declared)"
NO_EXCLUSIONS = "None declared"


def _emissions(**per_category) -> dict:
    fields = {
        "travel_tco2e": 0.0,
        "venue_energy_tco2e": 0.0,
        "accommodation_tco2e": 0.0,
        "catering_tco2e": 0.0,
        "materials_waste_tco2e": 0.0,
        "equipment_tco2e": 0.0,
        "swag_tco2e": 0.0,
        "digital_tco2e": 0.0,
    }
    fields.update(per_category)
    fields["total_tco2e"] = round(sum(fields.values()), 4)
    return fields


class TestCoverageMath:
    def test_every_category_actual_is_full_coverage(self):
        emissions = _emissions(travel_tco2e=6.0, catering_tco2e=4.0)
        assumptions = {"category_data_quality": {"travel": "actual", "catering": "actual"}}

        assert primary_data_coverage_pct(emissions, assumptions) == 100.0

    def test_share_is_weighted_by_tonnes_not_by_category_count(self):
        # 6 of 10 tCO2e measured, spread over one of two categories.
        emissions = _emissions(travel_tco2e=6.0, catering_tco2e=4.0)
        assumptions = {"category_data_quality": {"travel": "actual", "catering": "proxy"}}

        assert primary_data_coverage_pct(emissions, assumptions) == 60.0

    def test_waste_flag_maps_onto_the_materials_waste_total(self):
        # The assumptions key is "waste"; the emissions field is "materials_waste_tco2e".
        emissions = _emissions(materials_waste_tco2e=2.0, travel_tco2e=2.0)
        assumptions = {"category_data_quality": {"waste": "actual", "travel": "proxy"}}

        assert primary_data_coverage_pct(emissions, assumptions) == 50.0

    def test_not_applicable_categories_do_not_dilute_coverage(self):
        # A virtual event: the physical categories are gated off, carry zero tCO2e,
        # and must not count against the categories that were actually measured.
        emissions = _emissions(digital_tco2e=3.0)
        assumptions = {
            "category_data_quality": {
                "digital": "actual",
                "travel": "not applicable (virtual event)",
                "venue_energy": "not applicable (virtual event)",
                "accommodation": "not applicable (virtual event)",
                "catering": "not applicable (virtual event)",
                "waste": "not applicable (virtual event)",
                "equipment": "not provided",
                "swag": "not provided",
            }
        }

        assert primary_data_coverage_pct(emissions, assumptions) == 100.0

    def test_partial_categories_are_not_counted_as_primary(self):
        # "partial" travel is a mix of measured and proxy data — conservatively it
        # earns no primary-data credit.
        emissions = _emissions(travel_tco2e=8.0, catering_tco2e=2.0)
        assumptions = {"category_data_quality": {"travel": "partial", "catering": "actual"}}

        assert primary_data_coverage_pct(emissions, assumptions) == 20.0

    def test_no_flags_recorded_is_zero_coverage(self):
        assert primary_data_coverage_pct(_emissions(travel_tco2e=5.0), {}) == 0.0

    def test_zero_footprint_has_no_defined_coverage(self):
        assert primary_data_coverage_pct(_emissions(), {"category_data_quality": {}}) is None

    def test_malformed_flags_are_ignored_rather_than_trusted(self):
        emissions = _emissions(travel_tco2e=5.0)
        assert primary_data_coverage_pct(emissions, {"category_data_quality": "actual"}) is None


class TestSerializerExposure:
    def _row(self, *, assumptions: dict, input_payload: dict) -> ScenarioDB:
        return ScenarioDB(
            id="00000000-0000-0000-0000-000000000009",
            name="Coverage row",
            event_name="Coverage Event",
            location="Singapore",
            event_type="conference",
            attendees=100,
            event_days=1,
            mode="basic",
            travel_tco2e=6.0,
            venue_energy_tco2e=0.0,
            accommodation_tco2e=0.0,
            catering_tco2e=4.0,
            materials_waste_tco2e=0.0,
            equipment_tco2e=0.0,
            swag_tco2e=0.0,
            digital_tco2e=0.0,
            total_tco2e=10.0,
            per_attendee_tco2e=0.1,
            scope1_tco2e=0.0,
            scope2_tco2e=0.0,
            scope3_tco2e=10.0,
            data_quality="partly_primary",
            assumptions=assumptions,
            input_payload=input_payload,
            factors_snapshot={},
        )

    def test_coverage_and_exclusions_reach_the_wire_payload(self):
        row = self._row(
            assumptions={"category_data_quality": {"travel": "actual", "catering": "proxy"}},
            input_payload={"exclusions": "Attendee commuting and pre-event site visits."},
        )

        payload = serialize_scenario(row)

        assert payload["coverage_pct"] == 60.0
        assert payload["exclusions"] == "Attendee commuting and pre-event site visits."

    def test_absent_exclusions_serialize_as_none(self):
        row = self._row(assumptions={}, input_payload={})

        assert serialize_scenario(row)["exclusions"] is None


class TestScenarioApi:
    def test_exclusions_round_trip_through_create_and_fetch(self, client: TestClient):
        headers = register_user(client, email="exclusions-api@example.com")
        created = create_scenario(
            client, headers, exclusions="Excludes attendee commuting and marketing print."
        )

        assert created["exclusions"] == "Excludes attendee commuting and marketing print."
        assert created["coverage_pct"] is not None

        fetched = client.get(f"/api/scenarios/{created['scenario_id']}", headers=headers).json()
        assert fetched["exclusions"] == "Excludes attendee commuting and marketing print."
        assert fetched["coverage_pct"] == created["coverage_pct"]

    def test_over_long_exclusions_are_rejected(self):
        with pytest.raises(ValidationError):
            EventScenarioInput(name="Too long", attendees=10, exclusions="x" * 2001)

    def test_exclusions_is_optional(self):
        assert EventScenarioInput(name="No boundary note", attendees=10).exclusions is None


class TestExportRendering:
    DECLARED = "Excludes attendee commuting; excludes pre-event site visits."

    def _scenario_with_exclusions(self, client: TestClient, headers: dict) -> str:
        return create_scenario(client, headers, exclusions=self.DECLARED)["scenario_id"]

    def test_json_report_carries_the_declaration_and_coverage(self, client: TestClient):
        headers = register_user(client, email="excl-json@example.com")
        scenario_id = self._scenario_with_exclusions(client, headers)

        payload = client.get(f"/api/exports/scenarios/{scenario_id}.json", headers=headers).json()

        assert payload["exclusions"] == self.DECLARED
        assert payload["coverage_pct"] == payload["scenario"]["coverage_pct"]
        assert payload["coverage_pct"] is not None

    def test_json_report_says_none_declared_when_empty(self, client: TestClient):
        headers = register_user(client, email="excl-json-empty@example.com")
        scenario_id = create_scenario(client, headers)["scenario_id"]

        payload = client.get(f"/api/exports/scenarios/{scenario_id}.json", headers=headers).json()

        assert payload["exclusions"] == NO_EXCLUSIONS

    def test_csv_report_renders_the_declaration_and_coverage(self, client: TestClient):
        headers = register_user(client, email="excl-csv@example.com")
        scenario_id = self._scenario_with_exclusions(client, headers)

        content = client.get(f"/api/exports/scenarios/{scenario_id}.csv", headers=headers).text

        assert f"methodology,exclusions,{EXCLUSIONS_LABEL}," in content
        assert self.DECLARED in content
        assert "metadata,coverage_pct,Primary-Data Coverage," in content

    def test_csv_report_says_none_declared_when_empty(self, client: TestClient):
        headers = register_user(client, email="excl-csv-empty@example.com")
        scenario_id = create_scenario(client, headers)["scenario_id"]

        content = client.get(f"/api/exports/scenarios/{scenario_id}.csv", headers=headers).text

        assert f"methodology,exclusions,{EXCLUSIONS_LABEL},{NO_EXCLUSIONS}," in content

    def test_xlsx_report_renders_the_declaration_and_coverage(self, client: TestClient):
        headers = register_user(client, email="excl-xlsx@example.com")
        scenario_id = self._scenario_with_exclusions(client, headers)

        response = client.get(f"/api/exports/scenarios/{scenario_id}.xlsx", headers=headers)
        workbook = load_workbook(io.BytesIO(response.content))

        summary = list(workbook["Report Summary"].iter_rows(values_only=True))
        assert any(row[0] == "Primary-Data Coverage %" for row in summary)

        assumptions = list(workbook["Assumptions"].iter_rows(values_only=True))
        assert (EXCLUSIONS_LABEL, self.DECLARED) in [(row[0], row[1]) for row in assumptions]

    def test_xlsx_report_says_none_declared_when_empty(self, client: TestClient):
        headers = register_user(client, email="excl-xlsx-empty@example.com")
        scenario_id = create_scenario(client, headers)["scenario_id"]

        response = client.get(f"/api/exports/scenarios/{scenario_id}.xlsx", headers=headers)
        workbook = load_workbook(io.BytesIO(response.content))
        rows = [(row[0], row[1]) for row in workbook["Assumptions"].iter_rows(values_only=True)]

        assert (EXCLUSIONS_LABEL, NO_EXCLUSIONS) in rows

    @pytest.mark.parametrize("declared", [True, False])
    def test_pdf_report_renders_with_and_without_a_declaration(
        self, client: TestClient, declared: bool
    ):
        # The rendered bytes are not text-searchable, so the boundary wording is
        # asserted on the shared methodology block below; this covers the wiring.
        headers = register_user(client, email=f"excl-pdf-{declared}@example.com")
        overrides = {"exclusions": self.DECLARED} if declared else {}
        scenario_id = create_scenario(client, headers, **overrides)["scenario_id"]

        response = client.get(f"/api/exports/scenarios/{scenario_id}.pdf", headers=headers)

        assert response.status_code == 200
        assert response.content.startswith(b"%PDF")


class TestMethodologyBlock:
    """The label/value pairs the PDF renders under its summary table."""

    def _report(self, exclusions: str) -> ScenarioReportPayload:
        return ScenarioReportPayload(
            report_title="Report",
            scenario={"emissions": {}},
            scope_breakdown=ScopeBreakdown(),
            offset_portfolio=OffsetPortfolioSummary(
                total_purchased_tco2e=0, total_retired_tco2e=0, total_cost_usd=0
            ),
            compliance=ComplianceReport(
                overall_score_pct=0, checks=[], mandatory_frameworks=[], penalty_risk_usd=0
            ),
            exclusions=exclusions,
        )

    def test_declared_exclusions_are_stated_verbatim(self):
        block = dict(methodology_block(self._report("Excludes attendee commuting.")))

        assert block[EXCLUSIONS_LABEL] == "Excludes attendee commuting."

    def test_undeclared_boundary_is_stated_not_omitted(self):
        assert dict(methodology_block(self._report(NO_EXCLUSIONS)))[EXCLUSIONS_LABEL] == NO_EXCLUSIONS

    def test_boundary_sits_between_methodology_and_disclaimer(self):
        labels = [label for label, _ in methodology_block(self._report(NO_EXCLUSIONS))]

        assert labels == ["Methodology", EXCLUSIONS_LABEL, "Disclaimer"]
